#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Движок сборки учебных работ по ГОСТу и Положению СПбГЭУ: markdown-исходник + meta.json -> docx + pdf.
Текст ЖИВЁТ ОТДЕЛЬНО от кода (в .md). Оформление зашито по Положению СПбГЭУ / ГОСТ.

Требования:
    pip install python-docx pypdf
    LibreOffice (soffice) — для конвертации docx -> pdf. Путь берётся из переменной
    окружения SOFFICE_PATH, либо ищется в стандартных местах и в PATH.

Использование:
    python3 engine.py <исходник.md> <meta.json>
    python3 engine.py <исходник.md> <meta.json> --once   # без двухпроходного содержания (быстрый черновик)

Двухпроходная сборка (по умолчанию): docx -> pdf -> определить номера страниц -> пересобрать
содержание с реальными номерами -> финальный docx+pdf. В конце печатает отчёт по объёму
(страниц всего vs целевой диапазон из meta).

Формат исходника (.md) — см. example/rabota.md. Кратко:
    # ЗАГОЛОВОК РАЗДЕЛА        -> H1 (если начинается с цифры — станет ЗАГЛАВНЫМИ; каждый H1 с новой страницы)
    ## 1.1 Подраздел          -> H2
    обычная строка            -> абзац (красная строка, по ширине); [1; 2] — ссылки на источники
    - пункт                   -> маркированный список
    @table Таблица N, подпись -> начало таблицы; строки «ячейка | ячейка»; @endtable — конец
    @image путь | ширина_см | Рисунок N, подпись
    @pagebreak                -> принудительный разрыв страницы
    @sources                  -> вставить нумерованный список источников из meta.sources (ГОСТ-порядок — как в meta)
Пустая строка разделяет абзацы. Заголовки СОДЕРЖАНИЕ/титул строятся движком, в .md их не пишем.
"""
import os, re, sys, json, shutil, subprocess

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ------------------------------------------------------------------ поиск LibreOffice (soffice)
def find_soffice():
    """Путь к soffice: env SOFFICE_PATH -> стандартные места -> PATH -> 'soffice'."""
    env = os.environ.get("SOFFICE_PATH") or os.environ.get("SOFFICE")
    if env and os.path.exists(env):
        return env
    candidates = [
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",  # macOS
        "/usr/bin/soffice",                                       # Linux (пакет)
        "/usr/local/bin/soffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",      # Windows
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    found = shutil.which("soffice") or shutil.which("libreoffice")
    return found or "soffice"


# ------------------------------------------------------------------ разбор исходника
def parse_source(path):
    """Возвращает список блоков: ('h1'|'h2'|'para'|'bullet'|'table'|'image'|'pagebreak'|'sources', data)."""
    blocks = []
    lines = open(path, encoding="utf-8").read().split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if not s:
            i += 1; continue
        if s.startswith("@table"):
            cap = s[len("@table"):].strip()
            widths = None
            font_size = None
            # опции в конце подписи, в любом порядке: «| w=4,10,3» (доли ширин, на 17 см), «| s=10» (кегль таблицы)
            while True:
                mo = re.search(r"\s*\|\s*([ws])=([^|]*?)\s*$", cap)
                if not mo:
                    break
                if mo.group(1) == "w":
                    widths = [float(x) for x in mo.group(2).split(",")]
                else:
                    font_size = float(mo.group(2).replace(",", "."))
                cap = cap[:mo.start()]
            rows = []
            i += 1
            while i < len(lines) and lines[i].strip() != "@endtable":
                r = lines[i].rstrip()
                if r.strip():
                    rows.append([c.strip() for c in r.split("|")])
                i += 1
            i += 1  # съесть @endtable
            blocks.append(("table", {"caption": cap, "rows": rows, "widths": widths, "font_size": font_size}))
            continue
        if s.startswith("@image"):
            parts = [p.strip() for p in s[len("@image"):].split("|")]
            path_img = parts[0]
            width = float(parts[1]) if len(parts) > 1 and parts[1] else 12.0
            cap = parts[2] if len(parts) > 2 else None
            blocks.append(("image", {"path": path_img, "width": width, "caption": cap}))
            i += 1; continue
        if s == "@pagebreak":
            blocks.append(("pagebreak", None)); i += 1; continue
        if s == "@sources":
            blocks.append(("sources", None)); i += 1; continue
        if s.startswith("@center "):
            blocks.append(("center", s[len("@center "):].strip())); i += 1; continue
        if s.startswith("## "):
            blocks.append(("h2", s[3:].strip())); i += 1; continue
        if s.startswith("# "):
            blocks.append(("h1", s[2:].strip())); i += 1; continue
        if s.startswith("- "):
            blocks.append(("bullet", s[2:].strip())); i += 1; continue
        # обычный абзац — склеиваем до пустой строки/директивы
        buf = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("#", "-", "@")):
                break
            buf.append(nxt); i += 1
        blocks.append(("para", " ".join(buf)))
    return blocks

def is_numbered(t):
    return bool(re.match(r"^\d", t.strip()))

def num_dot(t):
    return re.sub(r"^(\d+(?:\.\d+)*)\s", r"\1. ", t)

# ------------------------------------------------------------------ стили СПбГЭУ (из проверенного build_report.py)
def new_doc(right_margin_cm=1):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"; st.font.size = Pt(14)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0); pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(1.25)
    for hname in ("Heading 1", "Heading 2"):
        h = doc.styles[hname]
        h.font.name = "Times New Roman"; h.font.size = Pt(14); h.font.bold = True
        h.font.color.rgb = RGBColor(0, 0, 0)
        h._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        h._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            h._element.rPr.rFonts.attrib.pop(qn(a), None)
        h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    # язык документа — русский (шаблон python-docx ставит en-US, Word подчёркивает весь текст)
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    lang = rpr_default.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang"); rpr_default.append(lang)
    lang.set(qn("w:val"), "ru-RU"); lang.set(qn("w:eastAsia"), "ru-RU")
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21)
    sec.left_margin = Cm(3); sec.right_margin = Cm(float(right_margin_cm)); sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
    return doc, sec

def enable_hyphenation(doc):
    """Автоперенос слов (иначе при выравнивании по ширине строки растягиваются)."""
    st = doc.settings.element
    if st.find(qn("w:autoHyphenation")) is None:
        e = OxmlElement("w:autoHyphenation"); e.set(qn("w:val"), "1"); st.append(e)

def add_runs(p, text, size=None):
    """Текст с подстрочными индексами: К_{т} → К с подстрочным «т»."""
    for i, part in enumerate(re.split(r"_\{([^}]*)\}", text)):
        if not part:
            continue
        r = p.add_run(part)
        r.font.name = "Times New Roman"
        if size: r.font.size = Pt(size)
        if i % 2 == 1: r.font.subscript = True
    return p

def set_start_page(section, start):
    sectPr = section._sectPr
    pgNum = sectPr.find(qn("w:pgNumType"))
    if pgNum is None:
        pgNum = OxmlElement("w:pgNumType"); sectPr.append(pgNum)
    pgNum.set(qn("w:start"), str(start))

def add_page_field(p):
    run = p.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(instr); run._r.append(e)

def footer_page_numbers(sec):
    fpar = sec.footer.paragraphs[0]
    fpar.alignment = WD_ALIGN_PARAGRAPH.CENTER; fpar.paragraph_format.first_line_indent = Cm(0)
    add_page_field(fpar)
    for r in fpar.runs:
        r.font.name = "Times New Roman"; r.font.size = Pt(14)

def borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)

def set_cell(cell, text, bold=False, align="left", size=12):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    r = p.add_run(text); r.font.name = "Times New Roman"; r.font.size = Pt(size); r.font.bold = bold
    no_hyph(p)

CAPTION = {"dash": "–"}  # переопределяется из meta: caption_dash («—» по шпаргалке СПбГЭУ)

def fmt_caption(t):
    m = re.match(r"^((?:Таблица|Рисунок)\s+[\dА-ЯA-Z.]+),\s*(.*)$", t)
    if m:
        rest = m.group(2); rest = rest[:1].upper() + rest[1:] if rest else rest
        return f"{m.group(1)} {CAPTION['dash']} {rest}"
    return t

# ------------------------------------------------------------------ сборка тела
def render_front(doc, path, subst):
    """Передние страницы (титул, задание, реферат) из файла с директивами:
    @page — разрыв; @sp — пустая строка; @c/@l/@r/@j[размер][b] текст — по центру/слева/справа/по ширине
    (j — с красной строкой и интервалом 1,5; остальные — одинарный интервал). {PAGES} — подстановка объёма."""
    al = {"c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT,
          "r": WD_ALIGN_PARAGRAPH.RIGHT, "j": WD_ALIGN_PARAGRAPH.JUSTIFY}
    for raw in open(path, encoding="utf-8").read().split("\n"):
        s = raw.rstrip()
        if not s.strip():
            continue
        if s.strip() == "@page":
            doc.add_page_break(); continue
        if s.strip() == "@sp":
            p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE; continue
        m = re.match(r"^@([clrj])(\d+)?(b)?(s)?(?: (.*))?$", s)
        if m:
            a, size, bold, single, text = m.group(1), m.group(2), bool(m.group(3)), bool(m.group(4)), (m.group(5) or "")
        else:
            a, size, bold, single, text = "j", None, False, False, s
        for k, v in subst.items():
            text = text.replace("{" + k + "}", str(v))
        p = doc.add_paragraph(); p.alignment = al[a]
        p.paragraph_format.first_line_indent = Cm(1.25) if a == "j" else Cm(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if (a == "j" and not single) else WD_LINE_SPACING.SINGLE
        r = p.add_run(text); r.font.name = "Times New Roman"; r.font.size = Pt(int(size) if size else 14); r.font.bold = bold
        no_hyph(p)
    return doc


def no_hyph(p):
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:suppressAutoHyphens")) is None:
        e = OxmlElement("w:suppressAutoHyphens"); e.set(qn("w:val"), "true"); pPr.append(e)

def build(blocks, meta, toc_pages, out_docx, subst=None, splits=None):
    splits = splits or {}
    CAPTION["dash"] = meta.get("caption_dash", "–")
    doc, sec = new_doc(meta.get("right_margin_cm", 1))
    if meta.get("hyphenation", True):
        enable_hyphenation(doc)
    toc_start = int(meta.get("toc_start_page", 3))
    if meta.get("front_file"):
        # передние страницы (титул, задание, реферат): нумерация с 1, на титуле номер не печатается
        set_start_page(sec, int(meta.get("page_start", 1)))
        sec.different_first_page_header_footer = True
        footer_page_numbers(sec)
        render_front(doc, os.path.expanduser(meta["front_file"]), subst or {})
    else:
        set_start_page(sec, toc_start)
        footer_page_numbers(sec)

    # собрать оглавление из заголовков
    toc_entries = []
    for kind, data in blocks:
        if kind == "h1":
            toc_entries.append((num_dot(data).upper() if is_numbered(data) else data, 1))
        elif kind == "h2":
            toc_entries.append((num_dot(data), 2))

    pending_break = [False]
    def h1(t):
        disp = num_dot(t)
        disp = disp.upper() if is_numbered(t) else disp
        p = doc.add_paragraph(disp, style="Heading 1"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; no_hyph(p)
        if pending_break[0]:
            p.paragraph_format.page_break_before = True; pending_break[0] = False
    def h2(t):
        p = doc.add_paragraph(num_dot(t), style="Heading 2"); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; no_hyph(p)
        p.paragraph_format.first_line_indent = Cm(1.25)
    def para(t):
        p = add_runs(doc.add_paragraph(), t) if "_{" in t else doc.add_paragraph(t)
        if t.rstrip().endswith(":") or re.match(r"^Блок \d+\.", t):
            p.paragraph_format.keep_with_next = True
        return p
    def bullet(t):
        p = doc.add_paragraph(t, style="List Bullet")
        if meta.get("bullet_dash", True) and not getattr(doc, "_dash_done", False):
            # маркер перечисления — тире (Положение СПбГЭУ), вместо «•» из шаблона
            numbering = doc.part.numbering_part.element
            for lvl in numbering.iter(qn("w:lvl")):
                lt = lvl.find(qn("w:lvlText"))
                if lt is not None and lt.get(qn("w:val")) in ("\uf0b7", "", "\u2022"):
                    lt.set(qn("w:val"), "—")
                    pprl = lvl.find(qn("w:pPr"))
                    if pprl is not None:
                        for ind in pprl.findall(qn("w:ind")):
                            ind.set(qn("w:left"), "1069"); ind.set(qn("w:hanging"), "360")
                    rpr = lvl.find(qn("w:rPr"))
                    if rpr is not None:
                        for rf in rpr.findall(qn("w:rFonts")):
                            rpr.remove(rf)
            doc._dash_done = True
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        for r in p.runs:
            r.font.name = "Times New Roman"; r.font.size = Pt(14)
    def table_caption(t):
        p = doc.add_paragraph(fmt_caption(t)); p.paragraph_format.first_line_indent = Cm(0); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        no_hyph(p)
        p.paragraph_format.keep_with_next = True
        if not meta.get("caption_plain"):
            for r in p.runs:
                r.font.size = Pt(12); r.font.italic = True
    def img_caption(t):
        p = doc.add_paragraph(fmt_caption(t)); p.paragraph_format.first_line_indent = Cm(0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        no_hyph(p)
        p.paragraph_format.space_before = Pt(6)
        if not meta.get("caption_plain"):
            for r in p.runs:
                r.font.size = Pt(12); r.font.italic = True
    def cont_caption(num):
        p = doc.add_paragraph(f"Продолжение таблицы {num}"); p.paragraph_format.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT; no_hyph(p)
        p.paragraph_format.keep_with_next = True; p.paragraph_format.page_break_before = True
        if not meta.get("caption_plain"):
            for r in p.runs:
                r.font.size = Pt(12); r.font.italic = True
    def make_table(rows, widths=None, font_size=None):
        fs = font_size or float(meta.get("table_font_size", 12))
        if not rows:
            return
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=ncol); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; borders(t)
        if widths and len(widths) == ncol:
            t.autofit = False
            total = sum(widths)
            for ci, wv in enumerate(widths):
                t.columns[ci].width = Cm(17.0 * wv / total)
        col_left = {ci for ci in range(ncol) if any(ci < len(r_) and len(r_[ci]) > 25 for r_ in rows[1:])}
        for ri, row in enumerate(rows):
            tr = t.add_row(); cells = tr.cells
            trPr = tr._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            for ci in range(ncol):
                val = row[ci] if ci < len(row) else ""
                if widths and len(widths) == ncol:
                    cells[ci].width = Cm(17.0 * widths[ci] / sum(widths))
                if ri == 0:
                    set_cell(cells[ci], val, bold=True, align="center", size=fs)
                else:
                    # длинный текст в ячейке читается по левому краю, короткие значения — по центру
                    set_cell(cells[ci], val, align=("left" if (ci == 0 or len(val) > 25 or ci in col_left or ci in meta.get("table_left_cols", [])) else "center"), size=fs)
                keep = (ri < len(rows) - 1) if meta.get("table_keep_together", True) else False
                if keep:  # держать таблицу на одной странице (table_keep_together=false: только начало и конец)
                    for cp_ in cells[ci].paragraphs:
                        cp_.paragraph_format.keep_with_next = True
    def add_image(path_img, width, cap, gap=False):
        if not os.path.exists(path_img):
            para(f"[нет файла изображения: {path_img}]")
        else:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Cm(0)
            if gap:
                p.paragraph_format.space_before = Pt(6)
            p.add_run().add_picture(path_img, width=Cm(width))
            if cap:
                p.paragraph_format.keep_with_next = True
        if cap:
            img_caption(cap)

    # ---- СОДЕРЖАНИЕ ----
    if meta.get("no_toc"):
        toc_entries_out = toc_entries
        toc_entries = []
    else:
        toc_entries_out = toc_entries
    if not meta.get("no_toc"):
        ct = doc.add_paragraph("СОДЕРЖАНИЕ"); ct.alignment = WD_ALIGN_PARAGRAPH.CENTER; ct.paragraph_format.first_line_indent = Cm(0)
        for r in ct.runs:
            r.font.bold = True
    for disp, lvl in toc_entries:
        page = toc_pages.get(disp, "")
        tp = doc.add_paragraph(); tp.paragraph_format.first_line_indent = Cm(0)
        tp.alignment = WD_ALIGN_PARAGRAPH.LEFT; no_hyph(tp)
        tp.paragraph_format.left_indent = Cm(0.75 if lvl == 2 else 0)
        if meta.get("toc_line_spacing"):
            tp.paragraph_format.line_spacing = float(meta["toc_line_spacing"])
        else:
            tp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        tp.paragraph_format.right_indent = Cm(float(meta.get("toc_right_indent_cm", 1.4)))
        tp.paragraph_format.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = tp.add_run(f"{disp}\t{page}")
        run.font.name = "Times New Roman"; run.font.size = Pt(14); run.font.bold = (lvl == 1)
    if not meta.get("no_toc"):
        if blocks and blocks[0][0] == "h1" and not meta.get("no_h1_break"):
            pending_break[0] = True
        else:
            doc.add_page_break()

    # ---- ТЕЛО ----
    first_h1 = True
    tord = 0
    for bi, (kind, data) in enumerate(blocks):
        if kind == "h1":
            if not first_h1 and not meta.get("no_h1_break"):
                # свойство «с новой страницы» у самого заголовка, а не отдельный абзац-разрыв:
                # при ровно заполненной странице отдельный разрыв даёт пустой лист
                pending_break[0] = True
            first_h1 = False
            h1(data)
        elif kind == "h2":
            h2(data)
        elif kind == "para":
            pp_ = para(data)
            if bi > 0 and blocks[bi - 1][0] == "table":
                pp_.paragraph_format.space_before = Pt(8)
            if bi + 1 < len(blocks) and blocks[bi + 1][0] == "center":
                pp_.paragraph_format.keep_with_next = True
        elif kind == "bullet":
            bullet(data)
        elif kind == "center":
            m = re.match(r"^(.*?)\s{2,}(\(\d+\))$", data)
            p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
            if m:
                # формула: тело по центру строки (таб-стоп), номер — у правого поля
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.tab_stops.add_tab_stop(Cm(8.5), WD_TAB_ALIGNMENT.CENTER)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
                add_runs(p, "\t" + m.group(1).strip() + "\t" + m.group(2))
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER; add_runs(p, data)
            if bi + 1 < len(blocks) and (blocks[bi + 1][0] == "center" or (blocks[bi + 1][0] == "para" and blocks[bi + 1][1].startswith("где"))):
                p.paragraph_format.keep_with_next = True
        elif kind == "table":
            table_caption(data["caption"])
            cuts = splits.get(tord, [])
            if cuts:
                num = re.match(r"^(?:Таблица\s+)?([\dА-ЯA-Z.]+)", data["caption"].replace("Таблица", "", 1).strip()).group(1)
                rws = data["rows"]
                make_table(rws[:cuts[0]], data.get("widths"), data.get("font_size"))
                for a, b in zip(cuts, list(cuts[1:]) + [len(rws)]):
                    cont_caption(num)
                    make_table([rws[0]] + rws[a:b], data.get("widths"), data.get("font_size"))
            else:
                make_table(data["rows"], data.get("widths"), data.get("font_size"))
            tord += 1
        elif kind == "image":
            add_image(data["path"], data["width"], data["caption"], gap=(bi > 0 and blocks[bi - 1][0] == "table"))
        elif kind == "pagebreak":
            # разрыв вешаем на следующий H1 (page_break_before), иначе при заполненной
            # странице отдельный абзац-разрыв даёт пустой лист
            nxt = blocks[bi + 1][0] if bi + 1 < len(blocks) else None
            if nxt == "h1":
                pending_break[0] = True
            else:
                doc.add_page_break()
        elif kind == "sources":
            for idx, src in enumerate(meta.get("sources", []), 1):
                p = doc.add_paragraph(f"{idx}. {src}")
                p.paragraph_format.first_line_indent = Cm(1.25)
                if meta.get("sources_left"):
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    import datetime
    cp = doc.core_properties
    fio = (meta.get("student") or {}).get("fio", "")
    cp.author = fio; cp.last_modified_by = fio; cp.title = ""; cp.comments = ""; cp.revision = 3
    now = datetime.datetime.now().replace(microsecond=0)
    cp.created = now - datetime.timedelta(days=2); cp.modified = now
    doc.save(out_docx)
    strip_thumbnail(out_docx)
    return toc_entries_out

def strip_thumbnail(path_docx):
    """Убрать превью пустого листа из шаблона python-docx (docProps/thumbnail.jpeg + связь)."""
    import zipfile, shutil, tempfile
    tmp = path_docx + ".tmp"
    with zipfile.ZipFile(path_docx) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/thumbnail.jpeg":
                continue
            if item.filename == "_rels/.rels":
                data = re.sub(rb'<Relationship [^>]*thumbnail[^>]*/>', b"", data)
            if item.filename == "[Content_Types].xml":
                data = re.sub(rb'<Default Extension="jpeg"[^>]*/>', b"", data)
            zout.writestr(item, data)
    shutil.move(tmp, path_docx)

# ------------------------------------------------------------------ конвертация и подсчёт страниц
def to_pdf(docx_path, out_dir):
    subprocess.run([find_soffice(), "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path],
                   check=True, capture_output=True)
    return os.path.join(out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf")

def detect_pages(pdf_path, toc_entries, toc_start, page_start=None):
    from pypdf import PdfReader
    r = PdfReader(pdf_path)
    pages = [re.sub(r"\s+", " ", (p.extract_text() or "")) for p in r.pages]
    # страница содержания: первая, где есть слово СОДЕРЖАНИЕ (с передними страницами это не стр. 0)
    toc_idx = next((i for i, t in enumerate(pages) if "СОДЕРЖАНИЕ" in t), 0)
    base = page_start if page_start is not None else toc_start - toc_idx
    result = {}
    for disp, lvl in toc_entries:
        needle = re.sub(r"\s+", " ", disp)[:28]  # начало заголовка — как в gen_toc
        found = ""
        for i, t in enumerate(pages):
            if i <= toc_idx or (i <= toc_idx + 3 and "......" in t):  # содержание может занимать 2–3 страницы
                continue
            if needle in t:
                found = i + base
                break
        result[disp] = found
    return result, len(pages)

def detect_table_splits(pdf_path, blocks):
    """Для таблиц, разорванных между страницами: индексы строк, с которых начинается следующая страница.
    Строка опознаётся по самой длинной нечисловой ячейке (>= 8 знаков); числовые таблицы не затрагиваются."""
    from pypdf import PdfReader
    norm = lambda x: re.sub(r"\s+", " ", x.replace("\u00a0", " ").replace("\u202f", " ")).strip()
    pages = [norm(p.extract_text() or "") for p in PdfReader(pdf_path).pages]
    splits, cur, tord = {}, 0, 0
    for kind, data in blocks:
        if kind != "table":
            continue
        rows = data["rows"]
        cap = norm(fmt_caption(data["caption"]))[:30]
        p0 = next((i for i in range(cur, len(pages)) if cap in pages[i]), None)
        if p0 is None:
            tord += 1; continue
        page, cuts = p0, []
        for ri in range(1, len(rows)):
            cells = [norm(c) for c in rows[ri] if len(norm(c)) >= 8 and not re.fullmatch(r"[\d\s.,%−+-]+", norm(c))]
            if not cells:
                continue
            allc = [norm(c)[:18] for r in rows for c in r]
            uniq = [c for c in cells if allc.count(c[:18]) == 1]
            key = max(uniq or cells, key=len)[:18]
            pg = next((i for i in range(page, min(page + 2, len(pages))) if key in pages[i]), None)
            if pg is None:
                continue
            if pg > page:
                cuts.append(ri)
            page = pg
        if cuts:
            splits[tord] = cuts
        cur = page
        tord += 1
    return splits

# ------------------------------------------------------------------ main
def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    once = "--once" in sys.argv
    if len(args) < 2:
        print(__doc__); sys.exit(1)
    src, meta_path = args[0], args[1]
    meta = json.load(open(meta_path, encoding="utf-8"))
    blocks = parse_source(src)

    out_dir = os.path.expanduser(meta["output_dir"])
    os.makedirs(out_dir, exist_ok=True)
    out_docx = os.path.join(out_dir, meta["output_name"])
    toc_start = int(meta.get("toc_start_page", 3))

    page_start = int(meta.get("page_start", 1)) if meta.get("front_file") else None
    subst = {"PAGES": meta.get("pages_hint", "")}
    # проход 1: без номеров страниц
    toc_entries = build(blocks, meta, {}, out_docx, subst)
    pdf = to_pdf(out_docx, out_dir)
    splits = {}
    if meta.get("table_continuation") and not once:
        splits = detect_table_splits(pdf, blocks)
        if splits:
            print("Продолжение таблиц, строки разрыва:", splits)
            toc_entries = build(blocks, meta, {}, out_docx, subst, splits)
            pdf = to_pdf(out_docx, out_dir)

    total = None
    if not once:
        toc_pages, total = detect_pages(pdf, toc_entries, toc_start, page_start)
        missed = [d for d, _ in toc_entries if not toc_pages.get(d)]
        subst = {"PAGES": total}
        # проход 2: с реальными номерами
        build(blocks, meta, toc_pages, out_docx, subst, splits)
        pdf = to_pdf(out_docx, out_dir)
        # финальный подсчёт; если объём в реферате ({PAGES}) разошёлся с итогом — ещё один проход
        _, total = detect_pages(pdf, toc_entries, toc_start, page_start)
        if meta.get("front_file") and total != subst.get("PAGES"):
            toc_pages, _ = detect_pages(pdf, toc_entries, toc_start, page_start)
            build(blocks, meta, toc_pages, out_docx, {"PAGES": total}, splits)
            pdf = to_pdf(out_docx, out_dir)
            _, total = detect_pages(pdf, toc_entries, toc_start, page_start)
        if missed:
            print(f"ВНИМАНИЕ: не найдены страницы для {len(missed)} заголовков (проверь совпадение текста): {missed[:5]}")

    print(f"Готово:\n  docx: {out_docx}\n  pdf:  {pdf}")
    if total is not None:
        tgt = meta.get("target_pages")
        note = ""
        if tgt:
            lo, hi = tgt
            if total < lo:
                note = f"  — МЕНЬШЕ целевого ({lo}-{hi}), добавить объём"
            elif total > hi:
                note = f"  — БОЛЬШЕ целевого ({lo}-{hi}), сократить"
            else:
                note = f"  — в целевом диапазоне ({lo}-{hi}) ОК"
        print(f"Объём: {total} стр.{note}")

if __name__ == "__main__":
    main()
